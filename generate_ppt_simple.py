from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# 创建演示文稿
prs = Presentation()

# 标题页
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "旅游票务系统"
subtitle = slide.placeholders[1]
subtitle.text = "毕业设计答辩\n\n王晓宇\n\n指导教师：XXX"

# 项目概述页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "项目概述"
content = slide.placeholders[1]
content.text = "项目背景：旅游业快速发展，线上票务需求增加\n\n项目目标：开发一个功能完整的旅游票务管理系统\n\n系统价值：简化购票流程，提高管理效率"

# 技术架构页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "技术架构"
content = slide.placeholders[1]
content.text = "技术栈：\n• 前端：HTML5, CSS3, JavaScript\n• 后端：Django, Python\n• 数据库：SQLite\n• API：天气查询API\n\nMVC架构：\n• 模型 (Models)：数据结构定义\n• 视图 (Views)：业务逻辑处理\n• 模板 (Templates)：页面展示"

# 系统功能模块页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "系统功能模块"
content = slide.placeholders[1]
content.text = "1. 用户系统\n   - 注册、登录、个人中心\n\n2. 景点管理\n   - 景点列表、详情、分类筛选\n\n3. 购票流程\n   - 门票选择、购物车、订单创建、支付\n\n4. 后台管理\n   - 平台管理员、景点管理员功能"

# 核心功能实现 - 用户系统
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "核心功能实现 - 用户系统"
content = slide.placeholders[1]
content.text = "实现细节：\n• 使用Django内置的认证系统\n• 自定义用户模型，添加角色字段\n• 实现三级权限控制：游客、景点管理员、平台管理员\n• 使用装饰器实现权限验证\n\n核心代码：\n@login_required\ndef personal_center(request):\n    user = request.user\n    # 处理用户信息更新"

# 核心功能实现 - 购票系统
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "核心功能实现 - 购票系统"
content = slide.placeholders[1]
content.text = "实现流程：\n1. 选择景点和门票类型\n2. 选择使用日期\n3. 检查库存\n4. 创建订单\n5. 支付\n6. 生成订单"

# 核心功能实现 - 后台管理
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "核心功能实现 - 后台管理"
content = slide.placeholders[1]
content.text = "平台管理员功能：\n• 用户管理：增删改查用户\n• 景点管理：审核、管理景点\n• 订单管理：查看所有订单\n• 评论管理：回复、删除评论\n\n景点管理员功能：\n• 景点管理：添加、编辑景点\n• 门票管理：设置门票类型和库存\n• 订单管理：查看景点订单\n• 评论管理：回复游客评论"

# 系统特色页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "系统特色"
content = slide.placeholders[1]
content.text = "1. 实时库存管理\n   - 按日期管理门票库存，确保数据准确性\n\n2. 天气查询集成\n   - 购票时查询景点天气，提升用户体验\n\n3. 多级权限管理\n   - 不同角色拥有不同权限，确保系统安全\n\n4. 响应式设计\n   - 适配不同设备，提供良好的用户体验"

# 数据库设计页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "数据库设计"
content = slide.placeholders[1]
content.text = "主要数据表：\n• User - 用户表\n• ScenicSpot - 景点表\n• TicketType - 门票类型表\n• Order - 订单表\n• Cart - 购物车表\n• ScenicSpotComment - 评论表\n• DateStock - 日期库存表"

# 系统演示页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "系统演示"
content = slide.placeholders[1]
content.text = "• 首页展示：轮播图、热门景点、最新资讯\n• 景点列表：分类筛选、搜索功能\n• 景点详情：景点信息、评论、购票入口\n• 购票流程：选择门票、日期、数量，创建订单\n• 个人中心：个人信息、订单管理、收藏\n• 后台管理：用户、景点、订单管理"

# 项目收获与展望页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "项目收获与展望"
content = slide.placeholders[1]
content.text = "项目收获：\n• 掌握Django框架的使用\n• 理解MVC架构设计模式\n• 学习数据库设计与优化\n• 提升问题解决能力\n\n未来展望：\n• 增加在线支付功能\n• 开发移动端应用\n• 优化系统性能\n• 添加更多旅游相关功能"

# 致谢页
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "谢谢聆听！"
subtitle = slide.placeholders[1]
subtitle.text = "欢迎提问"

# 保存PPT文件
prs.save("旅游票务系统毕业设计答辩.pptx")
print("PPT生成完成！")
