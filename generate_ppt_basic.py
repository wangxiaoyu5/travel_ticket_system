from pptx import Presentation

# 创建演示文稿
prs = Presentation()

# 标题页
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "旅游票务系统"
subtitle = slide.placeholders[1]
subtitle.text = "毕业设计答辩\n王晓宇\n指导教师：XXX"

# 项目概述页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "项目概述"
content = slide.placeholders[1]
content.text = "项目背景：旅游业快速发展，线上票务需求增加\n\n项目目标：开发一个功能完整的旅游票务管理系统\n\n系统价值：简化购票流程，提高管理效率"

# 技术架构页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "技术架构"
content = slide.placeholders[1]
content.text = "技术栈：\n- 前端：HTML5, CSS3, JavaScript\n- 后端：Django, Python\n- 数据库：SQLite\n- API：天气查询API\n\nMVC架构：\n- 模型 (Models)：数据结构定义\n- 视图 (Views)：业务逻辑处理\n- 模板 (Templates)：页面展示"

# 系统功能模块页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "系统功能模块"
content = slide.placeholders[1]
content.text = "1. 用户系统\n   - 注册、登录、个人中心\n\n2. 景点管理\n   - 景点列表、详情、分类筛选\n\n3. 购票流程\n   - 门票选择、购物车、订单创建、支付\n\n4. 后台管理\n   - 平台管理员、景点管理员功能"

# 核心功能实现页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "核心功能实现"
content = slide.placeholders[1]
content.text = "用户系统：\n- 使用Django内置认证系统\n- 三级权限控制\n- 装饰器实现权限验证\n\n购票系统：\n- 日期化库存管理\n- 实时库存检查\n- 订单生成与支付\n\n后台管理：\n- 平台管理员：管理所有用户、景点、订单\n- 景点管理员：管理所属景点信息和订单"

# 系统特色页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "系统特色"
content = slide.placeholders[1]
content.text = "1. 实时库存管理\n   - 按日期管理门票库存\n   - 实时显示库存状态\n\n2. 天气查询集成\n   - 购票时查询景点天气\n   - 根据天气提供游玩建议\n\n3. 多级权限管理\n   - 不同角色拥有不同权限\n   - 细粒度的权限控制\n\n4. 响应式设计\n   - 适配不同设备\n   - 移动端友好"

# 数据库设计页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "数据库设计"
content = slide.placeholders[1]
content.text = "主要数据表：\n- User - 用户表\n- ScenicSpot - 景点表\n- TicketType - 门票类型表\n- Order - 订单表\n- Cart - 购物车表\n- ScenicSpotComment - 评论表\n- DateStock - 日期库存表"

# 项目收获与展望页
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "项目收获与展望"
content = slide.placeholders[1]
content.text = "项目收获：\n- 掌握Django框架的使用\n- 理解MVC架构设计模式\n- 学习数据库设计与优化\n- 提升问题解决能力\n\n未来展望：\n- 增加在线支付功能\n- 开发移动端应用\n- 优化系统性能\n- 添加更多旅游相关功能"

# 致谢页
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "谢谢聆听！"
subtitle = slide.placeholders[1]
subtitle.text = "欢迎提问"

# 保存PPT文件
prs.save("旅游票务系统毕业设计答辩.pptx")
print("PPT生成完成！")
